from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from apps.workflow.services.data_transform import run_data_transform
from apps.workflow.services.demo_templates import BUILTIN_TEMPLATE_DIR
from apps.workflow.services.health_summary_sql import health_summary_queries
from apps.workflow.services.reports import parse_report_template, render_report

TEMPLATE_DIR = BUILTIN_TEMPLATE_DIR


def _metric(
    *,
    category: str,
    object_name: str,
    metric_name: str,
    value: float,
    unit: str,
    health_status: str,
    detail: str,
    collected_at: str,
) -> dict:
    return {
        "category": category,
        "object_type": "resource",
        "object_name": object_name,
        "dimensions": f"name={object_name}",
        "metric_name": metric_name,
        "value": value,
        "unit": unit,
        "warning_threshold": 80,
        "critical_threshold": 90,
        "health_status": health_status,
        "collected_at": collected_at,
        "detail": detail,
    }


def _host_result(
    *,
    name: str,
    ip: str,
    operating_system: str,
    conclusion: str,
    metrics: list[dict],
    collected_at: str,
    hostname: str | None = None,
) -> dict:
    critical = [item for item in metrics if item["health_status"] == "CRITICAL"]
    warning = [item for item in metrics if item["health_status"] == "WARNING"]
    normal = [item for item in metrics if item["health_status"] == "NORMAL"]
    return {
        "target": {
            "name": name,
            "ip": ip,
            "operating_system": operating_system,
        },
        "status": "SUCCESS",
        "result": {
            "collected_at": collected_at,
            "conclusion": conclusion,
            "host": {
                "hostname": hostname or name,
                "os_version": "Windows Server 2022" if operating_system == "windows" else "Ubuntu 22.04 LTS",
                "architecture": "x86_64",
                "uptime_hours": 100,
                "cpu_cores": 4,
                "memory_total_gb": 16,
                "top_cpu": "python=10%",
                "top_memory": "java=20%",
            },
            "metric_count": len(metrics),
            "critical_count": len(critical),
            "warning_count": len(warning),
            "normal_count": len(normal),
            "critical": critical,
            "warning": warning,
            "normal": normal,
            "metrics": metrics,
        },
        "error": None,
    }


HEALTH_REPORT_DATA = {
    "summary": {"total": 1, "succeeded": 1, "failed": 0},
    "data": [
        _host_result(
            name="job-web3",
            ip="10.10.90.120",
            operating_system="windows",
            conclusion="健康",
            collected_at="2026-09-18 10:30:00",
            metrics=[
                _metric(
                    category="计算",
                    object_name="CPU Total",
                    metric_name="usage_percent",
                    value=42.3,
                    unit="%",
                    health_status="NORMAL",
                    detail="CPU 使用率正常",
                    collected_at="2026-09-18 10:30:00",
                ),
                _metric(
                    category="磁盘",
                    object_name="C:",
                    metric_name="usage_percent",
                    value=73.1,
                    unit="%",
                    health_status="NORMAL",
                    detail="系统盘剩余空间充足",
                    collected_at="2026-09-18 10:30:00",
                ),
                _metric(
                    category="网络",
                    object_name="Ethernet0",
                    metric_name="throughput_mbps",
                    value=18.6,
                    unit="Mbps",
                    health_status="NORMAL",
                    detail="网卡接收速率正常",
                    collected_at="2026-09-18 10:30:00",
                ),
            ],
        )
    ],
}


MULTI_HOST_REPORT_DATA = {
    "summary": {"total": 2, "succeeded": 2, "failed": 0},
    "data": [
        _host_result(
            name="job-web3",
            ip="10.10.90.120",
            operating_system="windows",
            conclusion="健康",
            collected_at="2026-09-23 18:00:00",
            metrics=[
                _metric(
                    category="计算",
                    object_name="CPU Total",
                    metric_name="usage_percent",
                    value=12.0,
                    unit="%",
                    health_status="NORMAL",
                    detail="Windows CPU",
                    collected_at="2026-09-23 18:00:00",
                )
            ],
        ),
        _host_result(
            name="job-lab-linux-yum",
            ip="192.168.64.5",
            operating_system="linux",
            conclusion="需关注",
            collected_at="2026-09-23 18:00:01",
            metrics=[
                _metric(
                    category="磁盘",
                    object_name="/mnt/lima-cidata",
                    metric_name="usage_percent",
                    value=100.0,
                    unit="%",
                    health_status="CRITICAL",
                    detail="Linux 磁盘已满",
                    collected_at="2026-09-23 18:00:01",
                )
            ],
        ),
    ],
}


def _rendered_text(fmt: str, content: bytes) -> str:
    if fmt == "docx":
        document = Document(BytesIO(content))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        chunks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(chunks)
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    return "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_builtin_health_report_template_matches_job_output_contract_and_renders_multidimensional_metrics(fmt):
    template_path = TEMPLATE_DIR / f"health-inspection-example.{fmt}"
    content = template_path.read_bytes()

    parsed = parse_report_template(content, fmt)
    rendered = render_report(content, fmt, run_data_transform(HEALTH_REPORT_DATA, health_summary_queries(), 30))

    assert "data" in parsed.loops
    assert "data.result.metrics" in parsed.loops
    assert "data.result.critical" in parsed.loops
    assert "data.result.warning" in parsed.loops
    assert "metric_summary" in parsed.loops
    assert "summary.total" in parsed.placeholders
    assert "summary.critical_hosts" in parsed.placeholders
    assert any(item.endswith("target.ip") for item in parsed.placeholders)
    assert any(item.endswith("target.operating_system") for item in parsed.placeholders)
    assert any("display_name" in item for item in parsed.placeholders)
    template_text = _rendered_text(fmt, content)
    assert "{{" in template_text or "{%" in template_text
    assert "{d." not in template_text
    assert "Windows" not in template_text
    assert "主机健康巡检报告" in template_text
    assert "依据本次采集结果整理" in template_text
    assert "作业平台健康巡检报告" not in template_text
    if fmt == "docx":
        assert "主机一览" in template_text
        assert "指标项一览" in template_text
        assert "严重问题" in template_text
        assert "台主机里" in template_text
        assert "有严重项为红" in template_text
    else:
        assert "台主机里" not in template_text
        workbook = load_workbook(BytesIO(content), read_only=True)
        assert workbook.sheetnames == ["报告概览", "主机一览", "指标项一览", "严重问题", "告警问题", "检查明细"]
    text = _rendered_text(fmt, rendered)
    assert "10.10.90.120" in text
    assert "CPU Total" in text
    assert "Ethernet0" in text
    assert "job-web3" in text
    assert "42.3%" in text
    assert "≥80%" in text
    assert "≥90%" in text
    assert "{{" not in text
    assert HEALTH_REPORT_DATA["summary"] == {"total": 1, "succeeded": 1, "failed": 0}
    assert "metric_summary" not in HEALTH_REPORT_DATA


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_builtin_health_report_renders_failed_host_with_null_data(fmt):
    template_path = TEMPLATE_DIR / f"health-inspection-example.{fmt}"
    content = template_path.read_bytes()
    payload = run_data_transform(
        {
            "summary": {"total": 1, "succeeded": 0, "failed": 1},
            "data": [
                {
                    "target": {"name": "job-web3", "ip": "10.10.90.120", "operating_system": "windows"},
                    "status": "FAILED",
                    "result": None,
                    "error": "command timed out",
                }
            ],
        },
        health_summary_queries(),
        30,
    )
    rendered = render_report(content, fmt, payload)
    text = _rendered_text(fmt, rendered)
    assert payload["summary"]["total"] == 1
    assert payload["summary"]["succeeded"] == 0
    assert payload["summary"]["failed"] == 1
    assert "10.10.90.120" in text
    assert "未采集" in text
    assert "command timed out" not in text
    assert "{{" not in text


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_builtin_health_report_template_renders_multi_host_multi_os(fmt):
    template_path = TEMPLATE_DIR / f"health-inspection-example.{fmt}"
    content = template_path.read_bytes()

    rendered = render_report(content, fmt, run_data_transform(MULTI_HOST_REPORT_DATA, health_summary_queries(), 30))
    text = _rendered_text(fmt, rendered)

    assert "10.10.90.120" in text
    assert "192.168.64.5" in text
    assert "windows" in text
    assert "linux" in text
    assert "job-lab-linux-yum" in text
    assert "/mnt/lima-cidata" in text
    assert "100%" in text
    assert "{{" not in text
    if fmt == "docx":
        assert "2 台主机里 1 台有严重项，0 台只有告警，1 台健康，0 台没有采集到数据。" in text
    else:
        book = load_workbook(BytesIO(rendered))
        host_headers = [cell.value for cell in next(book["主机一览"].iter_rows(max_row=1))]
        metric_headers = [cell.value for cell in next(book["指标项一览"].iter_rows(max_row=1))]
        assert host_headers[:7] == ["主机", "IP", "系统", "严重项", "告警项", "正常项", "结论"]
        assert metric_headers[:8] == ["分类", "指标", "告警阈值", "严重阈值", "严重", "告警", "正常", "结论"]
        assert book["主机一览"].conditional_formatting._cf_rules
        assert book["报告概览"].sheet_view.view == "normal"
        host_rows = [[cell.value for cell in row] for row in book["主机一览"].iter_rows(max_col=7)]
        assert host_rows[1] == ["job-lab-linux-yum", "192.168.64.5", "linux", 1, 0, 0, "需关注"]
        assert host_rows[2] == ["job-web3", "10.10.90.120", "windows", 0, 0, 1, "健康"]
        assert book["主机一览"].max_row == 3
        metric_rows = [[cell.value for cell in row] for row in book["指标项一览"].iter_rows(max_col=8)]
        assert metric_rows[1][0] == "磁盘"
        assert metric_rows[1][4:8] == [1, 0, 0, "需关注"]
        assert metric_rows[2][0] == "计算"
        assert metric_rows[2][7] == "健康"
        assert book["报告概览"]["A4"].value == 2
        assert book["报告概览"]["D4"].value == 1
