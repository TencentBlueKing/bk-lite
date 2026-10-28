"""操作日志与登录日志导出超过上限时返回 400，不生成、不截断工作簿。"""

from io import BytesIO
from unittest.mock import patch

import pytest
from django.test import override_settings
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.system_mgmt.models import Group, OperationLog, User, UserLoginLog
from apps.system_mgmt.utils.audit_log_export import EXPORT_MAX_ROWS, configured_export_max_rows
from apps.system_mgmt.utils.group_filter_mixin import GroupFilterMixin
from apps.system_mgmt.utils.i18n import system_mgmt_message
from apps.system_mgmt.viewset.operation_log_viewset import OperationLogViewSet
from apps.system_mgmt.viewset.user_login_log_viewset import UserLoginLogViewSet

pytestmark = pytest.mark.django_db

VISIBLE_DOMAIN = "domain.com"
HIDDEN_DOMAIN = "other.test"


@pytest.fixture(autouse=True)
def scope_export_queryset_like_current_team(monkeypatch):
    """SQLite 不支持 JSON contains，无法执行 GroupFilterMixin 的组织过滤。

    这里用 domain 代替 current_team 范围，锁定导出上限统计的是 get_queryset()
    返回的查询集。组织外的行既不计入上限，也不进入工作簿。
    """

    def _visible(self):
        return self.queryset.filter(domain=VISIBLE_DOMAIN)

    monkeypatch.setattr(GroupFilterMixin, "get_queryset", _visible)


def test_export_max_rows_defaults_to_10000():
    assert EXPORT_MAX_ROWS == 10000
    assert configured_export_max_rows() == 10000


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS="2500")
def test_configured_export_max_rows_reads_settings():
    assert configured_export_max_rows() == 2500


@pytest.mark.parametrize("raw", [0, -3, "nope", "", None])
def test_configured_export_max_rows_rejects_invalid_values(raw):
    with override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=raw):
        assert configured_export_max_rows() == EXPORT_MAX_ROWS


def _admin():
    from apps.base.models import User as BaseUser

    admin = BaseUser.objects.create_user(
        username="export-cap-admin",
        password="pw",
        domain="domain.com",
        locale="en",
        email="export-cap-admin@x.com",
    )
    admin.is_superuser = True
    admin.save()
    return admin


def _team_and_user(username):
    group, _created = Group.objects.get_or_create(name="Default", parent_id=0)
    team_id = group.id
    User.objects.create(
        username=username,
        password="x",
        display_name=username,
        email=f"{username}@x.com",
        domain=VISIBLE_DOMAIN,
        group_list=[team_id],
    )
    return team_id


def _post_export(viewset, admin, team_id, body):
    request = APIRequestFactory().post("/export_excel/", body, format="json")
    request.COOKIES["current_team"] = str(team_id)
    force_authenticate(request, user=admin)
    return viewset.as_view({"post": "export_excel"})(request)


def _assert_rejected(response, workbook, count, limit=1):
    workbook.assert_not_called()
    assert response.status_code == 400
    assert response.get("Content-Disposition") is None
    if hasattr(response, "render"):
        response.render()
    assert "spreadsheetml" not in response["Content-Type"]
    assert response.content[:2] != b"PK"
    expected = system_mgmt_message("en", "export.too_many_rows", limit=limit, count=count)
    assert response.data["message"] == expected


def _sheet(response):
    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    return load_workbook(BytesIO(response.content)).active


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=1)
def test_operation_log_export_over_limit_returns_400_without_workbook():
    admin = _admin()
    team_id = _team_and_user("oplog-user")
    for index in range(2):
        OperationLog.objects.create(
            username="oplog-user",
            domain=VISIBLE_DOMAIN,
            source_ip="127.0.0.1",
            app="cmdb",
            action_type="create",
            summary=f"op-row-{index}",
        )
    OperationLog.objects.create(
        username="oplog-user",
        domain=HIDDEN_DOMAIN,
        source_ip="127.0.0.1",
        app="cmdb",
        action_type="create",
        summary="op-hidden",
    )

    with patch("apps.system_mgmt.viewset.operation_log_viewset.Workbook") as workbook:
        response = _post_export(OperationLogViewSet, admin, team_id, {})

    _assert_rejected(response, workbook, count=2)
    assert b"op-row-0" not in response.content
    assert b"op-hidden" not in response.content
    assert OperationLog.objects.filter(username="oplog-user").count() == 3


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=2)
def test_operation_log_export_at_limit_keeps_every_row():
    admin = _admin()
    team_id = _team_and_user("oplog-full")
    for index in range(2):
        OperationLog.objects.create(
            username="oplog-full",
            domain=VISIBLE_DOMAIN,
            source_ip="127.0.0.1",
            app="cmdb",
            action_type="create",
            summary=f"kept-{index}",
        )
    OperationLog.objects.create(
        username="oplog-full",
        domain=HIDDEN_DOMAIN,
        source_ip="127.0.0.1",
        app="cmdb",
        action_type="create",
        summary="kept-hidden",
    )

    response = _post_export(OperationLogViewSet, admin, team_id, {})
    sheet = _sheet(response)
    assert sheet.max_row == 3
    assert {sheet.cell(row=2, column=7).value, sheet.cell(row=3, column=7).value} == {"kept-0", "kept-1"}
    assert "kept-hidden" not in {sheet.cell(row=row, column=7).value for row in range(2, sheet.max_row + 1)}


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=1)
def test_operation_log_selected_ids_within_limit_exports_only_those_rows():
    admin = _admin()
    team_id = _team_and_user("oplog-pick")
    chosen = OperationLog.objects.create(
        username="oplog-pick",
        domain=VISIBLE_DOMAIN,
        source_ip="127.0.0.1",
        app="cmdb",
        action_type="create",
        summary="chosen",
    )
    OperationLog.objects.create(
        username="oplog-pick",
        domain=VISIBLE_DOMAIN,
        source_ip="127.0.0.1",
        app="cmdb",
        action_type="update",
        summary="left-out",
    )

    response = _post_export(OperationLogViewSet, admin, team_id, {"selected_ids": [chosen.id]})
    sheet = _sheet(response)
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=7).value == "chosen"


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=1)
def test_operation_log_selected_ids_over_limit_returns_400():
    admin = _admin()
    team_id = _team_and_user("oplog-many")
    rows = [
        OperationLog.objects.create(
            username="oplog-many",
            domain=VISIBLE_DOMAIN,
            source_ip="127.0.0.1",
            app="cmdb",
            action_type="create",
            summary=f"id-row-{index}",
        )
        for index in range(2)
    ]

    with patch("apps.system_mgmt.viewset.operation_log_viewset.Workbook") as workbook:
        response = _post_export(
            OperationLogViewSet,
            admin,
            team_id,
            {"selected_ids": [row.id for row in rows]},
        )

    _assert_rejected(response, workbook, count=2)


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=1)
def test_operation_log_narrow_filter_exports_when_filtered_count_is_within_limit():
    admin = _admin()
    team_id = _team_and_user("alice-op")
    _team_and_user("bob-op")
    OperationLog.objects.create(
        username="alice-op",
        domain=VISIBLE_DOMAIN,
        source_ip="127.0.0.1",
        app="cmdb",
        action_type="create",
        summary="alice-only",
    )
    OperationLog.objects.create(
        username="bob-op",
        domain=VISIBLE_DOMAIN,
        source_ip="127.0.0.1",
        app="cmdb",
        action_type="create",
        summary="bob-only",
    )

    response = _post_export(OperationLogViewSet, admin, team_id, {"username": "alice-op"})
    sheet = _sheet(response)
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=7).value == "alice-only"


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=1)
def test_login_log_export_over_limit_returns_400_without_workbook():
    admin = _admin()
    team_id = _team_and_user("login-user")
    for index in range(2):
        UserLoginLog.objects.create(
            username="login-user",
            domain=VISIBLE_DOMAIN,
            source_ip="127.0.0.1",
            status=UserLoginLog.STATUS_SUCCESS,
            location=f"login-row-{index}",
        )
    UserLoginLog.objects.create(
        username="login-user",
        domain=HIDDEN_DOMAIN,
        source_ip="127.0.0.1",
        status=UserLoginLog.STATUS_SUCCESS,
        location="login-hidden",
    )

    with patch("apps.system_mgmt.viewset.user_login_log_viewset.Workbook") as workbook:
        response = _post_export(UserLoginLogViewSet, admin, team_id, {})

    _assert_rejected(response, workbook, count=2)
    assert b"login-row-0" not in response.content


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=2)
def test_login_log_export_at_limit_keeps_every_row():
    admin = _admin()
    team_id = _team_and_user("login-full")
    for index in range(2):
        UserLoginLog.objects.create(
            username="login-full",
            domain=VISIBLE_DOMAIN,
            source_ip="127.0.0.1",
            status=UserLoginLog.STATUS_FAILED,
            location=f"kept-login-{index}",
        )
    UserLoginLog.objects.create(
        username="login-full",
        domain=HIDDEN_DOMAIN,
        source_ip="127.0.0.1",
        status=UserLoginLog.STATUS_SUCCESS,
        location="kept-login-hidden",
    )

    response = _post_export(UserLoginLogViewSet, admin, team_id, {})
    sheet = _sheet(response)
    assert sheet.max_row == 3
    values = {sheet.cell(row=2, column=5).value, sheet.cell(row=3, column=5).value}
    assert values == {"kept-login-0", "kept-login-1"}


@override_settings(AUDIT_LOG_EXPORT_MAX_ROWS=1)
def test_login_log_selected_ids_over_limit_returns_400_and_within_limit_exports():
    admin = _admin()
    team_id = _team_and_user("login-pick")
    chosen = UserLoginLog.objects.create(
        username="login-pick",
        domain=VISIBLE_DOMAIN,
        source_ip="10.0.0.8",
        status=UserLoginLog.STATUS_SUCCESS,
        location="chosen-login",
    )
    other = UserLoginLog.objects.create(
        username="login-pick",
        domain=VISIBLE_DOMAIN,
        source_ip="10.0.0.9",
        status=UserLoginLog.STATUS_FAILED,
        location="other-login",
    )

    with patch("apps.system_mgmt.viewset.user_login_log_viewset.Workbook") as workbook:
        rejected = _post_export(
            UserLoginLogViewSet,
            admin,
            team_id,
            {"selected_ids": [chosen.id, other.id]},
        )
    _assert_rejected(rejected, workbook, count=2)

    exported = _post_export(UserLoginLogViewSet, admin, team_id, {"selected_ids": [chosen.id]})
    sheet = _sheet(exported)
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=5).value == "chosen-login"
