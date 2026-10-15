"""migrate / Django setup 不得加载 python-docx、reportlab；附件生成仍可用。"""

import ast
from pathlib import Path

import pytest

from apps.opspilot.services.workflow_attachment_service import _iter_markdown_blocks, build_attachment_bytes, normalize_attachment_file_type

pytestmark = pytest.mark.unit

_OPSPILOT_ROOT = Path(__file__).resolve().parents[1]


def _top_level_imported_modules(path: Path) -> set[str]:
    tree = ast.parse(path.read_text(encoding="utf-8"))
    names: set[str] = set()
    for node in tree.body:
        if isinstance(node, ast.Import):
            names.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            names.add(node.module)
    return names


def test_nats_api_does_not_import_chat_flow_factory_at_module_level():
    names = _top_level_imported_modules(_OPSPILOT_ROOT / "nats_api.py")
    assert "apps.opspilot.utils.chat_flow_utils.engine.factory" not in names


def test_workflow_attachment_service_defers_docx_and_reportlab():
    names = _top_level_imported_modules(_OPSPILOT_ROOT / "services" / "workflow_attachment_service.py")
    assert "docx" not in names
    assert "openpyxl" not in names
    assert "pptx" not in names
    assert not any(name == "reportlab" or name.startswith("reportlab.") for name in names)


def test_build_attachment_bytes_md_keeps_utf8():
    assert build_attachment_bytes("# 报告", "md", title="日报") == "# 报告".encode("utf-8")


def test_build_attachment_bytes_txt_keeps_plain_text():
    payload = build_attachment_bytes("广州介绍", "txt")
    assert payload == "广州介绍".encode("utf-8")
    _, filename, mime_type = normalize_attachment_file_type("txt", "广州介绍.txt")
    assert filename == "广州介绍.txt"
    assert mime_type == "text/plain"


def test_build_attachment_bytes_docx_renders_markdown_heading_and_table():
    from io import BytesIO

    from docx import Document

    payload = build_attachment_bytes(
        "# 主机磁盘资源爆满风险评估报告\n\n## 报告信息\n\n| 项目 | 内容 |\n| --- | --- |\n| 报告日期 | 2026-10-12 |\n",
        "docx",
        title="主机磁盘资源爆满风险评估报告",
    )
    document = Document(BytesIO(payload))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "主机磁盘资源爆满风险评估报告" in texts
    assert not any(text.startswith("#") for text in texts)
    assert not any("|" in text for text in texts)
    assert document.tables[0].rows[1].cells[0].text == "报告日期"
    assert document.tables[0].rows[1].cells[1].text == "2026-10-12"
    assert document.paragraphs[0].runs[0].font.name == "微软雅黑"


def test_build_attachment_bytes_pptx_is_a_slide_deck():
    from io import BytesIO

    from pptx import Presentation

    payload = build_attachment_bytes("# 广州旅游\n\n- 珠江\n- 早茶\n", "pptx", title="广州旅游")
    assert payload[:2] == b"PK"
    deck = Presentation(BytesIO(payload))
    assert len(deck.slides) == 1
    assert deck.slides[0].shapes.title.text == "广州旅游"
    assert "珠江" in deck.slides[0].placeholders[1].text_frame.text
    assert deck.slides[0].placeholders[1].text_frame.paragraphs[0].runs[0].font.name == "微软雅黑"


def test_build_attachment_bytes_docx_is_office_zip():
    payload = build_attachment_bytes("第一行\n第二行", "docx", title="巡检")
    assert payload[:2] == b"PK"
    assert len(payload) > 1000


def test_build_attachment_bytes_pdf_has_header():
    payload = build_attachment_bytes("hello <tag> & value", "pdf", title="Inspection")
    assert payload.startswith(b"%PDF")
    assert len(payload) > 200


def test_build_attachment_bytes_csv_has_utf8_bom():
    payload = build_attachment_bytes("name,count\nweb,2", "csv")
    assert payload.startswith(b"\xef\xbb\xbf")
    assert payload.decode("utf-8-sig") == "name,count\nweb,2"


@pytest.mark.django_db
def test_refresh_attachment_download_urls_resigns_expired_history_link(monkeypatch):
    import json

    from django.core.files.base import ContentFile
    from django.core.signing import SignatureExpired

    from apps.opspilot.models import WorkflowAttachmentAsset
    from apps.opspilot.services import workflow_attachment_service as attachment_service

    asset = WorkflowAttachmentAsset.objects.create(
        execution_id="exec-history",
        attachment_id="att-history",
        filename="子网地址使用率.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file=ContentFile(b"xlsx", name="subnet.xlsx"),
    )
    stale = attachment_service.build_signed_attachment_download_url(asset)
    token = stale.rsplit("/download/", 1)[1].rstrip("/")
    stored = json.dumps(
        [{"type": "TOOL_CALL_RESULT", "content": json.dumps({"file_url": stale, "filename": asset.filename})}],
        ensure_ascii=False,
    )
    monkeypatch.setattr(attachment_service, "workflow_attachment_download_max_age", lambda: -1)
    with pytest.raises(SignatureExpired):
        attachment_service.resolve_signed_attachment_token(token)

    refreshed = attachment_service.refresh_attachment_download_urls(stored)
    fresh_token = refreshed.split("/download/", 1)[1].split("/", 1)[0]
    monkeypatch.setattr(attachment_service, "workflow_attachment_download_max_age", lambda: 24 * 60 * 60)
    resolved = attachment_service.resolve_signed_attachment_token(fresh_token)
    assert resolved is not None
    assert resolved.id == asset.id


def test_markdown_paragraph_stops_before_list_or_rule_without_blank_line():
    blocks = list(_iter_markdown_blocks("介绍\n- 项\n1. 序号\n---\n结尾"))
    assert blocks == [
        ("paragraph", "介绍"),
        ("ul", ["项"]),
        ("ol", ["序号"]),
        ("hr",),
        ("paragraph", "结尾"),
    ]
    assert list(_iter_markdown_blocks("介绍\n继续")) == [("paragraph", "介绍 继续")]


def test_build_attachment_bytes_pptx_splits_each_h1_into_a_slide():
    from io import BytesIO

    from pptx import Presentation

    payload = build_attachment_bytes("# 第一页\n\n要点\n\n# 第二页\n\n收尾", "pptx", title="第一页")
    deck = Presentation(BytesIO(payload))
    titles = [slide.shapes.title.text for slide in deck.slides]
    assert len(deck.slides) >= 2
    assert titles == ["第一页", "第二页"]


def test_normalize_attachment_file_type_rewrites_alias_extension():
    normalized_type, filename, mime_type = normalize_attachment_file_type("ppt", "deck.ppt")
    assert normalized_type == "ppt"
    assert filename == "deck.pptx"
    assert mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    _, text_name, text_mime = normalize_attachment_file_type("text", "notes.text")
    assert text_name == "notes.txt"
    assert text_mime == "text/plain"

    _, word_name, word_mime = normalize_attachment_file_type("word", "memo.doc")
    assert word_name == "memo.docx"
    assert word_mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    _, kept_name, _ = normalize_attachment_file_type("pptx", "deck.pptx")
    assert kept_name == "deck.pptx"


def _pdf_plain_text(payload: bytes) -> str | None:
    try:
        import fitz
    except ImportError:
        return None
    document = fitz.open(stream=payload, filetype="pdf")
    try:
        return "".join(page.get_text() for page in document)
    finally:
        document.close()


def test_build_attachment_bytes_pdf_renders_chinese_markdown_table():
    from apps.opspilot.services.workflow_attachment_service import _resolve_pdf_font

    content = "| 项目 | 内容 |\n| --- | --- |\n| 城市 | 广州 |\n| 备注 | A & B <C> |\n"
    font_name = _resolve_pdf_font()
    first = build_attachment_bytes(content, "pdf", title="广州报告")
    second = build_attachment_bytes(content, "pdf", title="广州报告")
    for payload in (first, second):
        assert payload.startswith(b"%PDF")
        assert len(payload) > 200
        # Helvetica 没有中文轮廓，fitz 会抽出占位符；有中文字体时必须能读回单元格。
        if font_name == "Helvetica":
            continue
        extracted = _pdf_plain_text(payload)
        if not extracted:
            continue
        assert "广州" in extracted
        assert "城市" in extracted
        assert "项目" in extracted
        assert "A & B" in extracted


def test_build_attachment_bytes_xlsx_keeps_header_row():
    from io import BytesIO

    from openpyxl import load_workbook

    payload = build_attachment_bytes("name,count\nweb,2", "xlsx")
    sheet = load_workbook(BytesIO(payload)).active
    assert [cell.value for cell in next(sheet.iter_rows(max_row=1))] == ["name", "count"]
    assert [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))] == ["web", "2"]
