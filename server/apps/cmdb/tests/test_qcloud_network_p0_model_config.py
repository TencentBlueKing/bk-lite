"""腾讯云 TKE / VPN / NAT P0 模型种子。"""
import os

import openpyxl
import pytest

pytestmark = pytest.mark.unit

XLSX = os.path.join(os.path.dirname(__file__), "..", "support-files", "model_config.xlsx")

MODELS = {
    "qcloud_tke_cluster": ("腾讯云TKE集群", "cc-k8s-cluster_K8S集群"),
    "qcloud_tke_node": ("腾讯云TKE节点", "cc-cloud-server_云服务器"),
    "qcloud_vpn_gateway": ("腾讯云VPN网关", "cc-cloud_云"),
    "qcloud_customer_gateway": ("腾讯云对端网关", "cc-cloud_云"),
    "qcloud_vpn_connection": ("腾讯云VPN通道", "cc-cloud_云"),
    "qcloud_nat_gateway": ("腾讯云NAT网关", "cc-cloud_云"),
}

INT_FIELDS = {
    "qcloud_tke_cluster": {"node_num", "master_node_num"},
    "qcloud_vpn_gateway": {"bandwidth_mbps", "max_connection"},
    "qcloud_nat_gateway": {"bandwidth_mbps", "max_concurrent"},
}


def _records(sheet):
    headers = [cell.value for cell in sheet[2]]
    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=3, values_only=True)
        if any(value is not None and str(value).strip() != "" for value in values)
    ]


def test_p0_models_belong_to_qcloud_with_reused_icons():
    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    models = {row["model_id"]: row for row in _records(workbook["models"])}
    for model_id, (model_name, icn) in MODELS.items():
        assert models[model_id]["classification_id"] == "qcloud"
        assert models[model_id]["model_name"] == model_name
        assert models[model_id]["icn"] == icn
        assert models[model_id]["app_topo_layer"] == "none"
        assert f"attr-{model_id}" in workbook.sheetnames
        assert f"asso-{model_id}" in workbook.sheetnames
    workbook.close()


def test_p0_attrs_and_association_pages():
    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    for model_id, int_fields in INT_FIELDS.items():
        attrs = {row["attr_id"]: row for row in _records(workbook[f"attr-{model_id}"])}
        assert attrs["resource_name"]["attr_type"] == "str"
        assert attrs["resource_id"]["attr_type"] == "str"
        assert attrs["tag"]["attr_type"] == "tag"
        assert attrs["create_time"]["attr_type"] == "str"
        for field in int_fields:
            assert attrs[field]["attr_type"] == "int"
        assert "PreShareKey" not in attrs
        assert "presharekey" not in attrs

    node_attrs = {row["attr_id"] for row in _records(workbook["attr-qcloud_tke_node"])}
    assert {"cluster_id", "ip_addr", "instance_role", "node_pool_id"} <= node_attrs

    expected = {
        ("qcloud_tke_cluster", "qcloud", "belong", "n:1"),
        ("qcloud_tke_node", "qcloud", "belong", "n:1"),
        ("qcloud_tke_node", "qcloud_tke_cluster", "belong", "n:1"),
        ("qcloud_tke_node", "qcloud_cvm", "connect", "n:1"),
        ("qcloud_vpn_gateway", "qcloud", "belong", "n:1"),
        ("qcloud_customer_gateway", "qcloud", "belong", "n:1"),
        ("qcloud_vpn_connection", "qcloud", "belong", "n:1"),
        ("qcloud_vpn_connection", "qcloud_vpn_gateway", "connect", "n:1"),
        ("qcloud_vpn_connection", "qcloud_customer_gateway", "connect", "n:1"),
        ("qcloud_nat_gateway", "qcloud", "belong", "n:1"),
        ("qcloud_nat_gateway", "qcloud_eip", "connect", "n:n"),
    }
    actual = set()
    for model_id in MODELS:
        for row in _records(workbook[f"asso-{model_id}"]):
            actual.add((row["src_model_id"], row["dst_model_id"], row["asst_id"], row["mapping"]))
    assert expected <= actual
    workbook.close()
